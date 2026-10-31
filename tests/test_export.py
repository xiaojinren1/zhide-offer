"""D exports: isolated temporary accounts, actual DOCX parsing, no model/network."""
import json
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from fastapi.testclient import TestClient
from docx import Document
from api.export import MIME, MIME_PDF
from core import accounts
from main import app

HEADER = {"X-Requested-With": "XMLHttpRequest"}
FIXTURE = Path(__file__).parent / "fixtures" / "role-d" / "resume.md"


class ExportContract(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory(prefix="zhide-export-")
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name)
        override = patch.object(accounts, "DATABASE_PATH", self.root / "accounts.sqlite3")
        override.start()
        self.addCleanup(override.stop)
        self.client = TestClient(app, headers=HEADER)
        self.addCleanup(self.client.close)
        self.payload = {"markdown": FIXTURE.read_text(encoding="utf-8-sig"), "name": "林小禾（虚构）"}

    def register(self):
        self.assertEqual(self.client.post("/api/auth/register", json={
            "username": "export_demo", "password": "demo-only-12345",
        }).status_code, 200)

    def test_anonymous_and_missing_csrf_are_rejected(self):
        self.assertEqual(self.client.post("/api/export/docx", json=self.payload).status_code, 401)
        self.register()
        self.client.headers.pop("X-Requested-With")
        self.assertEqual(self.client.post("/api/export/docx", json=self.payload).status_code, 403)

    def test_actual_docx_has_chinese_headings_lists_and_no_saved_file(self):
        self.register()
        response = self.client.post("/api/export/docx", json=self.payload)
        self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else "")
        self.assertEqual(response.headers["content-type"], MIME)
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertIn("filename*=UTF-8''", response.headers["content-disposition"])
        doc = Document(BytesIO(response.content))
        paragraphs = doc.paragraphs
        text = "\n".join(p.text for p in paragraphs)
        self.assertIn("林小禾", text)
        self.assertIn("校园活动报名工具", text)
        self.assertIn("需求清单", text)
        self.assertNotIn("**", text)
        self.assertTrue(any(p.style.name == "Heading 1" for p in paragraphs))
        self.assertTrue(any(p.style.name == "List Bullet" for p in paragraphs))
        self.assertTrue(any(r.bold and "产品协作成员" in r.text for p in paragraphs for r in p.runs))
        self.assertEqual(list(self.root.glob("*.docx")), [])

    def test_actual_pdf_has_pdf_signature_and_no_saved_file(self):
        self.register()
        response = self.client.post("/api/export/pdf", json=self.payload)
        self.assertEqual(response.status_code, 200, response.text if response.status_code != 200 else "")
        self.assertEqual(response.headers["content-type"], MIME_PDF)
        self.assertEqual(response.headers["cache-control"], "no-store")
        self.assertIn("filename*=UTF-8''", response.headers["content-disposition"])
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertGreater(len(response.content), 1000)
        self.assertEqual(list(self.root.glob("*.pdf")), [])

    def test_pdf_reuses_auth_and_input_contract(self):
        self.assertEqual(self.client.post("/api/export/pdf", json=self.payload).status_code, 401)
        self.register()
        for value in ("", "x" * 100001, "bad\x00text", "bad\ud800text", "bad\ufffftext"):
            with self.subTest(value=repr(value[:12])):
                response = self.client.post("/api/export/pdf", content=json.dumps({"markdown": value}), headers={"Content-Type": "application/json"})
                self.assertEqual(response.status_code, 422)

    def test_blank_oversized_and_invalid_xml_are_rejected(self):
        self.register()
        for value in ("", " \n\t", "x" * 100001, "bad\x00text", "bad\ud800text", "bad\ufffftext"):
            with self.subTest(value=repr(value[:12])):
                response = self.client.post("/api/export/docx", content=json.dumps({"markdown": value}), headers={"Content-Type": "application/json"})
                self.assertEqual(response.status_code, 422)

    def test_export_does_not_fetch_links_or_create_external_relationships(self):
        self.register()
        markdown = "# 演示\n[链接文字](https://example.invalid/private)\n![图片](file:///etc/passwd)\n1. 一步\n---\n<script>alert(1)</script>"
        response = self.client.post("/api/export/docx", json={"markdown": markdown, "name": '../x\r\nHeader: y'})
        self.assertEqual(response.status_code, 200)
        with ZipFile(BytesIO(response.content)) as archive:
            relationships = archive.read("word/_rels/document.xml.rels").decode()
            self.assertNotIn('TargetMode="External"', relationships)
            self.assertNotIn("example.invalid", relationships)
        text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
        self.assertIn("链接文字", text)
        self.assertIn("<script>alert(1)</script>", text)
        self.assertNotIn("\r", response.headers["content-disposition"])
        self.assertNotIn("\n", response.headers["content-disposition"])

    def test_expired_session_cannot_export(self):
        self.register()
        self.assertEqual(self.client.post("/api/auth/logout").status_code, 200)
        self.assertEqual(self.client.post("/api/export/docx", json=self.payload).status_code, 401)


if __name__ == "__main__":
    unittest.main()
