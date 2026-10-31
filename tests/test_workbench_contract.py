import sys
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from fastapi.testclient import TestClient
from main import app
from engine.resume_parser import profile_from_dict
from core.accounts import require_user


class AuthenticatedContract(unittest.TestCase):
    """业务测试隔离认证依赖；真实 Cookie 和用户归属由 test_accounts 覆盖。"""
    def setUp(self):
        previous = dict(app.dependency_overrides)
        app.dependency_overrides[require_user] = lambda: {'id': 'contract-user', 'username': 'contract_user'}
        self.addCleanup(lambda: (app.dependency_overrides.clear(), app.dependency_overrides.update(previous)))

class WorkbenchContract(AuthenticatedContract):
    def setUp(self):
        super().setUp()
        configured = patch("api.resume.is_configured", return_value=True)
        configured.start()
        self.addCleanup(configured.stop)
        # 新增规划/复核阶段也必须 Mock，契约测试不能访问真实模型。
        def tailoring_json(_system, user, **_kwargs):
            import json
            payload = json.loads(user)
            if 'jd' in payload:
                return {'requirements': [{'quote': unit, 'priority': 'core', 'evidence_ids': ['e1.actions.1']} for unit in payload['units']]}
            if 'review' in payload:
                return {'markdown': '# 演示同学\n整理报名需求'}
            return {'items': [{'id': row['id'], 'quote': ''} for row in payload['plan']['requirements']], 'unsupported_claims': []}
        planner = patch('engine.tailoring.chat_json', side_effect=tailoring_json)
        planner.start()
        self.addCleanup(planner.stop)
        self.client = TestClient(app, headers={'X-Requested-With': 'XMLHttpRequest'})
        self.profile = {'name':'演示同学','experiences':[{'name':'报名流程','actions':['整理需求'],'results':[],'skills':[]}], 'skills':[], 'education':[], 'user_confirmed':True}

    def test_confirmation_preserved(self):
        self.assertFalse(profile_from_dict({}).user_confirmed)
        self.assertTrue(profile_from_dict(self.profile).user_confirmed)

    def test_unconfirmed_never_calls_model(self):
        with patch('api.resume.chat_with_tools') as model:
            for value in (False, None, 'true'):
                response = self.client.post('/api/resume/generate',json={'jd_text':'产品经理','resume_profile':dict(self.profile,user_confirmed=value)})
                self.assertEqual(response.status_code,422)
            model.assert_not_called()

    def test_missing_inputs_rejected(self):
        for jd,profile in [('',self.profile),('产品经理',dict(self.profile,experiences=[]))]:
            response=self.client.post('/api/resume/generate',json={'jd_text':jd,'resume_profile':profile})
            self.assertEqual(response.status_code,422)

    def test_parse_and_generate_contract(self):
        with patch('api.resume.is_configured',return_value=True), patch('engine.resume_parser.chat_json',return_value=self.profile):
            data=self.client.post('/api/resume/parse-text',json={'text':'演示同学整理报名需求'}).json()
            self.assertTrue(data['ok'])
            self.assertFalse(data['profile']['user_confirmed'])
        with patch('api.resume.chat_with_tools',return_value={'content':'# 演示同学\n整理报名需求','rounds':1,'tool_calls':0}) as model:
            data=self.client.post('/api/resume/generate',json={'jd_text':'产品经理','resume_profile':self.profile}).json()
            self.assertTrue(data['ok'])
            self.assertIn('演示同学',model.call_args.args[1])

    def test_interrupted_draft_not_delivered(self):
        with patch('api.resume.chat_with_tools',return_value={'content':'未完成草稿','deadline_hit':True}):
            response=self.client.post('/api/resume/generate',json={'jd_text':'产品经理','resume_profile':self.profile})
            self.assertEqual(response.status_code,504)
            self.assertNotIn('markdown',response.json())

    def test_empty_and_provider_failure(self):
        for result in ('',RuntimeError('test upstream failure')):
            args={'side_effect':result} if isinstance(result,Exception) else {'return_value':{'content':result}}
            with patch('api.resume.chat_with_tools',**args):
                response=self.client.post('/api/resume/generate',json={'jd_text':'产品经理','resume_profile':self.profile})
                self.assertTrue(response.status_code==502 or response.json().get('ok') is False)

class FileAndStorageContract(AuthenticatedContract):
    def test_unquantified_source_does_not_require_invented_metrics(self):
        from engine.models import ResumeProfile, Experience
        from engine.quality_check import check_markdown
        profile = ResumeProfile(experiences=[Experience(name='需求梳理', actions=['收集反馈'], results=['形成原型'])])
        report = check_markdown('# 演示同学\n收集反馈，形成原型', profile)
        self.assertEqual(report['high_count'], 0)
        findings = [i for i in report['issues'] if i['type'] == 'no_quantification']
        self.assertTrue(findings)
        self.assertTrue(all('不新增数字或量级' in i['fix'] for i in findings))

    def test_docx_tables_preserved(self):
        import io
        from docx import Document
        from engine.resume_parser import extract_text
        doc=Document()
        doc.add_paragraph('演示经历')
        doc.add_table(rows=1,cols=2).rows[0].cells[0].text='表格中的项目'
        stream=io.BytesIO(); doc.save(stream)
        self.assertIn('表格中的项目',extract_text('resume.docx',stream.getvalue()))

    def test_confirmed_dates_are_known_numbers(self):
        from engine.models import ResumeProfile, Experience
        from engine.quality_check import untraceable_numbers
        profile=ResumeProfile(education=['2022 入学'],experiences=[Experience(period='2024.03-2024.09')])
        self.assertEqual(untraceable_numbers('2022 入学；项目2024.03-2024.09',profile),[])
        self.assertEqual(untraceable_numbers('增长99%',profile),['99'])

    def test_profile_ids_do_not_collide(self):
        from api.profiles import _new_id
        self.assertNotEqual(_new_id('demo'),_new_id('demo'))

    def test_unconfigured_parse_explicit_error(self):
        with patch('api.resume.is_configured',return_value=False), patch('api.resume.parse_resume') as model:
            response=TestClient(app, headers={'X-Requested-With': 'XMLHttpRequest'}).post('/api/resume/parse-text',json={'text':'示例经历'})
            self.assertEqual(response.status_code,503)
            model.assert_not_called()

if __name__=='__main__': unittest.main()
